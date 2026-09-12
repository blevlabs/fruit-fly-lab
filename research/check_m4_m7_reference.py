"""Read-only MaleCNS gamma1 circuit/source audit; no neural model or runtime edits.

Locally: python check_m4_m7_reference.py --check
With prepared data: python research/check_m4_m7_reference.py --prepared runtime/malecns-v1
Prints JSON to stdout. Redirect it locally to retain an evidence receipt.
"""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import statistics

import numpy as np
import pandas as pd


def incoming(ids, indptr, indices, counts, target, allowed):
    """Extract real presynaptic edges; CSR rows are postsynaptic, not sources."""
    row = np.searchsorted(ids, target)
    if row == len(ids) or ids[row] != target:
        raise ValueError(f'Target {target} absent from graph')
    bounds = slice(indptr[row], indptr[row + 1])
    sources, weights = ids[indices[bounds]], counts[bounds]
    return [[int(pre), int(weight)] for pre, weight in zip(sources, weights)
            if int(pre) in allowed]


def check():
    # Actual incoming edges to cell 20 are 10->20 (2 contacts), 40->20 (7).
    ids = np.array([10, 20, 40])
    ptr = np.array([0, 0, 2, 3])
    idx, weights = np.array([0, 2, 1]), np.array([2, 7, 3])
    assert incoming(ids, ptr, idx, weights, 20, {10, 40}) == [[10, 2], [40, 7]]
    assert incoming(ids, ptr, idx, weights, 10, {20}) == []
    try:
        incoming(ids, ptr, idx, weights, 30, {10})
    except ValueError:
        pass
    else:
        raise AssertionError('Absent target silently accepted')


def audit(directory):
    neurons = pd.read_feather(directory / 'neurons.feather')
    with np.load(directory / 'connectome.npz') as graph:
        ids, ptr, idx, counts, signs = [graph[k] for k in
            ('neuron_ids', 'indptr', 'indices', 'synapse_count', 'fast_current_sign')]
    assert neurons.bodyId.is_unique and np.all(np.diff(ids) > 0)
    assert np.array_equal(ids, neurons.bodyId.to_numpy())
    assert len(ptr) == len(ids) + 1 and ptr[0] == 0 and ptr[-1] == len(idx) == len(counts)
    assert np.all(counts > 0) and np.all(idx >= 0) and np.all(idx < len(ids))
    gamma = neurons.loc[neurons.type.fillna('').str.startswith('KCg')]
    targets = neurons.loc[neurons.type.isin(['MBON11', 'PPL101'])]
    assert len(targets) == 4 and targets.type.value_counts().to_dict() == {'MBON11': 2, 'PPL101': 2}
    gamma_ids = set(map(int, gamma.bodyId))
    dan_ids = set(map(int, targets.loc[targets.type.eq('PPL101'), 'bodyId']))
    all_selected = gamma_ids | set(map(int, targets.bodyId))
    edges = []
    for target in sorted(all_selected):
        # KC->MBON and PPL101->KC/MBON are retrieved independently; no triangle inferred.
        allowed = gamma_ids | dan_ids if target in set(targets.bodyId) else dan_ids
        edges.extend({'pre_bodyId': pre, 'post_bodyId': target, 'synapse_count': count}
                     for pre, count in incoming(ids, ptr, idx, counts, target, allowed))
    unresolved = signs == 0
    out_unresolved = unresolved[idx]
    fields = ['bodyId', 'type', 'instance', 'somaSide', 'hemibrainType', 'flywireType',
              'consensus_nt', 'ground_truth', 'predicted_nt', 'receptorType']
    summary = []
    for target in targets.itertuples():
        group = [e for e in edges if e['post_bodyId'] == target.bodyId and e['pre_bodyId'] in gamma_ids]
        summary.append({'bodyId': int(target.bodyId), 'instance': target.instance,
                        'incoming_gamma_kc_edges': len(group),
                        'incoming_gamma_kc_contacts': sum(e['synapse_count'] for e in group)})
    paths = [directory / 'neurons.feather', directory / 'connectome.npz']
    return {'claim': 'exact-ID anatomical reference extraction only',
            'implementation': 'PASS: source retention, CSR direction and cell/edge extraction',
            'numerical_physiology': 'NOT_RUN', 'biological_validation': 'OPEN', 'learning_capability': 'NOT_RUN',
            'preparation': 'adult male MaleCNS v1.0 anatomical specimen; not the physiology animals',
            'source_sha256': {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in paths},
            'neuron_count': len(ids), 'edge_count': len(counts), 'synapse_contacts': int(counts.sum()),
            'unresolved_neurons': int(unresolved.sum()),
            'unresolved_outgoing_edges': int(out_unresolved.sum()),
            'unresolved_outgoing_contacts': int(counts[out_unresolved].sum()),
            'gamma_kc_count': len(gamma), 'reference_cells': json.loads(targets[fields].to_json(orient='records')),
            'target_summary': summary, 'gamma_kc_ids': sorted(gamma_ids), 'selected_edges': edges,
            'limitations': ['Type correspondence does not identify recorded animals or odor-selective KCs.',
                'Somatic side is retained as annotation, not used to fabricate lateralized contacts.',
                'Neuron-pair contacts lack compartment coordinates, receptor identity and chemical kinetics.',
                'PPL101 contacts do not establish a pointwise dopamine eligibility mask or volume transmission.',
                'No weights, currents, receptor assignments, learned state or behavioral decoder were changed.']}


def cibarial_reference(path):
    """Extract published observations, without inferring tau/C or simulating spikes."""
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    groups, expected = ['MN12V', 'MN12D', 'MN11V', 'MN11D'], [10, 8, 8, 13]
    rows = []
    specs = [('EDF5a', 'threshold', 'mV', groups, expected),
             ('EDF5b', 'resting_potential', 'mV', groups, expected),
             ('EDF5c', 'input_resistance', 'Mohm', ['MN12D', 'MN11V'], [6, 6]),
             ('EDF5f', 'first_spike_delay_source_convention', 'ms', ['MN12D-MN11V'], [6])]
    for sheet, metric, unit, names, sizes in specs:
        for column, (name, count) in enumerate(zip(names, sizes), 1):
            cells = [book[sheet].cell(r, column) for r in range(2, book[sheet].max_row + 1)]
            values = [float(c.value) for c in cells if c.value is not None]
            assert len(values) == count and np.isfinite(values).all()
            rows.append({'cell_class': name, 'metric': metric, 'unit': unit,
                         'sheet': sheet, 'source_cells': [c.coordinate for c in cells if c.value is not None],
                         'observations': values, 'n': len(values), 'mean': statistics.mean(values),
                         'sample_sd': statistics.stdev(values)})
    book.close()
    resistance = [r['observations'] for r in rows if r['metric'] == 'input_resistance']
    differences = [right - left for left, right in zip(*resistance)]
    return {'status': 'PUBLISHED_STATIC_REFERENCE_EXTRACTED_NOT_A_DYNAMICAL_FIT',
            'source': 'https://www.nature.com/articles/s41593-026-02412-y/figures/12',
            'preparation': 'In vivo adult Drosophila feeding recordings for EDF5a/b; same-pair current injection for EDF5c-f. Sex/age/temperature are not identified per observation.',
            'genotype_authority': 'Supplementary Table 1: EDF5a/b use MN12V (VT031562/GMR75F02), MN12D/11V (VT050240/GMR10E04), MN11D (NP0534); EDF5c-f use VT050240/GMR10E04 with mCD8-GFP.',
            'workbook_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
            'groups_and_units_authority': 'Extended Data Figure 5 axes and caption, visually verified',
            'observations': rows,
            'paired_MN11V_minus_MN12D_resistance_Mohm': differences,
            'paired_difference_mean_Mohm': statistics.mean(differences),
            'limitations': ['Source sheets a/b lack column labels; figure axis fixes the retained order.',
                'Paired recording labels do not provide individual fly IDs or cross-panel pairing.',
                'No capacitance, membrane tau or raw current/voltage ramp traces in this workbook.',
                'Recording-site thresholds are not silently substituted for effective LIF parameters.',
                'No biological fit, runtime assignment, spike sequence or feeding capability is claimed.']}


def synapse_audit(directory):
    """Reconcile bounded public contacts with the pinned prepared graph subset."""
    reference = json.loads((directory.parent / 'circuit-audit.json').read_text())
    gamma = set(reference['gamma_kc_ids'])
    mbons, dans = {10704, 11402}, {11327, 11900}
    expected = {(e['pre_bodyId'], e['post_bodyId']): e['synapse_count']
                for e in reference['selected_edges']
                if (e['pre_bodyId'] in gamma and e['post_bodyId'] in mbons)
                or (e['pre_bodyId'] in dans and e['post_bodyId'] in gamma | mbons)}
    manifest = json.loads((directory / 'gamma1-source-retrieval.json').read_text())
    observed, results = Counter(), []
    for source in manifest['sources']:
        if not source['path'].endswith('-synapses.json'):
            continue
        path = directory / Path(source['path']).name
        raw = path.read_bytes()
        assert hashlib.sha256(raw).hexdigest() == source['sha256']
        data = json.loads(raw)
        assert len(data['data']) == source['rows']
        rois = Counter()
        for row in data['data']:
            r = dict(zip(data['columns'], row))
            pair = (r['body_pre'], r['body_post'])
            if pair in expected:
                observed[pair] += 1
                rois.update(set(r['rois_pre']) & set(r['rois_post']) &
                            {'g1(L)', 'g1(R)', 'PED(L)', 'PED(R)'})
            assert all(np.isfinite(r[k]) for k in
                       ('x_pre', 'y_pre', 'z_pre', 'x_post', 'y_post', 'z_post'))
        results.append({'file': path.name, 'contacts': source['rows'],
                        'prepared_subset_same_ROI_contact_counts': dict(rois)})
    assert len(results) == 4 and dict(observed) == expected
    return {'status': 'PASS: measured contact counts reproduce prepared anatomical subset',
            'overlap_neuron_pairs': len(expected), 'total_retrieved_contact_pairs': sum(x['contacts'] for x in results),
            'coordinate_units': '8 nm voxel, neuPrint male-cns:v1.0 metadata', 'files': results,
            'physiology_or_learning_validation': 'NOT_RUN',
            'limitations': ['ROI membership does not establish receptor action or a plasticity eligibility mask.',
                           'All measured contacts retained; no soma-side or axon/dendrite filter imposed.']}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--prepared', type=Path)
    parser.add_argument('--check', action='store_true')
    parser.add_argument('--cibarial', type=Path)
    parser.add_argument('--synapses', type=Path)
    args = parser.parse_args()
    check()
    if sum(x is not None for x in [args.prepared, args.cibarial, args.synapses]) > 1:
        parser.error('Choose one of --prepared, --cibarial or --synapses')
    if args.synapses is not None:
        print(json.dumps(synapse_audit(args.synapses), indent=2, allow_nan=False))
    elif args.cibarial is not None:
        print(json.dumps(cibarial_reference(args.cibarial), indent=2, allow_nan=False))
    elif args.prepared is not None:
        print(json.dumps(audit(args.prepared), indent=2, allow_nan=False))
    elif args.check:
        print('PASS: actual CSR direction and missing-target rejection; no neural simulation')
    else:
        parser.error('--prepared, --cibarial, --synapses or --check is required')
